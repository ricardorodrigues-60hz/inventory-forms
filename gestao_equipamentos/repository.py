import io
import re
import traceback
from datetime import datetime
from typing import Any, Dict, List, Optional

import core
from core import to_thread, acquire_cimed_connection, release_cimed_connection
from modules.gestao_equipamentos.schemas import (
    EquipamentoCreate, EquipamentoUpdate, NovoEquipamentoImport, ParMerge
)

class EquipamentoRepository:
    """Repositório responsável por todo o acesso a dados, consultas SQL e regras de enriquecimento/agrupamento de equipamentos."""

    @staticmethod
    async def execute_dml_cimed(sql: str, params: Optional[Dict[str, Any]] = None) -> int:
        def run():
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                cur.execute(sql, params or {})
                rowcount = cur.rowcount
                conn.commit()
                cur.close()
                return rowcount
            finally:
                release_cimed_connection(conn)

        return await to_thread(run)

    @staticmethod
    async def execute_batch_insert_cimed(items: List[Dict[str, Any]]) -> int:
        def run():
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql = """
                    INSERT INTO ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS (
                        cd_patrimonio,
                        cd_setor,
                        local_especifico,
                        tipo_maquina,
                        observacao,
                        cd_serie,
                        cd_mac,
                        nr_ip,
                        nm_hostname,
                        registrado_por_id,
                        registrado_em,
                        sn_inserido_manual
                    ) VALUES (
                        :cd_patrimonio,
                        :cd_setor,
                        :local_especifico,
                        :tipo_maquina,
                        :observacao,
                        :cd_serie,
                        :cd_mac,
                        :nr_ip,
                        :nm_hostname,
                        :registrado_por_id,
                        SYSDATE,
                        NVL(:sn_inserido_manual, 'N')
                    )
                """
                cur.executemany(sql, items)
                rowcount = cur.rowcount
                conn.commit()
                cur.close()
                return rowcount
            finally:
                release_cimed_connection(conn)

        return await to_thread(run)

    @staticmethod
    async def enriquecer_dados_equipamento(cd_patrimonio: str) -> Dict[str, Optional[str]]:
        """
        Realiza o enriquecimento automático de dados do equipamento a partir do cd_patrimonio.
        1. Busca nr_serie (deq_numserie as cd_serie) no Oracle CIMED (contrato 1012700).
        2. Se encontrado, busca nr_ip, nm_hostname e cd_mac no OCS MariaDB.
        """
        resultado: Dict[str, Optional[str]] = {
            "cd_serie": None,
            "nr_ip": None,
            "nm_hostname": None,
            "cd_mac": None,
        }

        patrimonio_limpo = str(cd_patrimonio or "").strip().lstrip("0")
        if not patrimonio_limpo:
            return resultado

        sql_serie = """
            SELECT a.deq_numserie AS cd_serie
            FROM cimed.man_detalheequipamento a
            INNER JOIN cimed.man_contratoentrega b ON a.cen_ky = b.cen_ky
            INNER JOIN cimed.man_contrato c ON b.con_ky = c.con_ky
            WHERE c.con_ky = 1012700
              AND a.deq_idcontrato = :cd_patrimonio
        """
        row_serie = await core.fetch_one_cimed(sql_serie, {"cd_patrimonio": patrimonio_limpo})
        cd_serie = (row_serie or {}).get("cd_serie")
        if not cd_serie:
            return resultado

        resultado["cd_serie"] = str(cd_serie).strip()

        sql_ocs = """
            SELECT h.NAME AS nm_hostname, h.IPADDR AS nr_ip, n.MACADDR AS cd_mac
            FROM hardware h
            INNER JOIN networks n ON h.ID = n.HARDWARE_ID
            INNER JOIN bios b ON h.ID = b.HARDWARE_ID
            WHERE b.SSN = %s
        """
        try:
            row_rede = await core.fetch_one_ocs(sql_ocs, (resultado["cd_serie"],))
            if row_rede:
                resultado["nr_ip"] = row_rede.get("nr_ip") or None
                resultado["nm_hostname"] = row_rede.get("nm_hostname") or None
                resultado["cd_mac"] = row_rede.get("cd_mac") or None
        except Exception as e:
            print(f"[Equipamentos] OCS MariaDB inacessível para cd_serie={resultado['cd_serie']}: {e}")

        return resultado

    @staticmethod
    def construir_filtros_equipamentos(
        q: Optional[str] = None,
        cd_setor: Optional[str] = None,
        registrado_por_id: Optional[str] = None,
        tipo_maquina: Optional[str] = None,
        dt_inicio: Optional[str] = None,
        dt_fim: Optional[str] = None,
        cd_serie: Optional[str] = None,
        nr_ip: Optional[str] = None,
        nm_hostname: Optional[str] = None,
        status_incompleto: Optional[str] = None,
    ) -> tuple[str, Dict[str, Any]]:
        where_clauses = ["1=1"]
        params: Dict[str, Any] = {}

        if q and q.strip():
            term = f"%{q.strip()}%"
            where_clauses.append("""
                (UPPER(e.cd_patrimonio) LIKE UPPER(:q)
                 OR UPPER(e.cd_setor) LIKE UPPER(:q)
                 OR UPPER(s.nm_setor) LIKE UPPER(:q)
                 OR UPPER(e.local_especifico) LIKE UPPER(:q)
                 OR UPPER(e.tipo_maquina) LIKE UPPER(:q)
                 OR UPPER(e.observacao) LIKE UPPER(:q)
                 OR UPPER(e.cd_serie) LIKE UPPER(:q)
                 OR UPPER(e.cd_mac) LIKE UPPER(:q)
                 OR UPPER(e.nr_ip) LIKE UPPER(:q)
                 OR UPPER(e.nm_hostname) LIKE UPPER(:q)
                 OR UPPER(e.registrado_por_id) LIKE UPPER(:q))
            """)
            params["q"] = term

        if cd_setor and cd_setor.strip():
            where_clauses.append("e.cd_setor = :cd_setor")
            params["cd_setor"] = cd_setor.strip()

        if registrado_por_id and registrado_por_id.strip():
            where_clauses.append("UPPER(e.registrado_por_id) = UPPER(:registrado_por_id)")
            params["registrado_por_id"] = registrado_por_id.strip()

        if tipo_maquina and tipo_maquina.strip() and tipo_maquina.strip().upper() != "TODOS":
            where_clauses.append("UPPER(e.tipo_maquina) = UPPER(:tipo_maquina)")
            params["tipo_maquina"] = tipo_maquina.strip().upper()

        if dt_inicio and dt_inicio.strip():
            val_inicio = dt_inicio.strip()
            if len(val_inicio) == 10:
                val_inicio += " 00:00:00"
            where_clauses.append("e.registrado_em >= TO_DATE(:dt_inicio, 'YYYY-MM-DD HH24:MI:SS')")
            params["dt_inicio"] = val_inicio

        if dt_fim and dt_fim.strip():
            val_fim = dt_fim.strip()
            if len(val_fim) == 10:
                val_fim += " 23:59:59"
            where_clauses.append("e.registrado_em <= TO_DATE(:dt_fim, 'YYYY-MM-DD HH24:MI:SS')")
            params["dt_fim"] = val_fim

        if cd_serie and cd_serie.strip():
            where_clauses.append("UPPER(e.cd_serie) LIKE UPPER(:cd_serie)")
            params["cd_serie"] = f"%{cd_serie.strip()}%"

        if nr_ip and nr_ip.strip():
            where_clauses.append("UPPER(e.nr_ip) LIKE UPPER(:nr_ip)")
            params["nr_ip"] = f"%{nr_ip.strip()}%"

        if nm_hostname and nm_hostname.strip():
            where_clauses.append("UPPER(e.nm_hostname) LIKE UPPER(:nm_hostname)")
            params["nm_hostname"] = f"%{nm_hostname.strip()}%"

        if status_incompleto and status_incompleto.strip():
            st = status_incompleto.strip().lower()
            if st == "sem_ip":
                where_clauses.append("(e.nr_ip IS NULL OR TRIM(e.nr_ip) = '' OR TRIM(e.nr_ip) = '-')")
            elif st == "sem_hostname":
                where_clauses.append("(e.nm_hostname IS NULL OR TRIM(e.nm_hostname) = '' OR TRIM(e.nm_hostname) = '-')")
            elif st == "sem_serie":
                where_clauses.append("(e.cd_serie IS NULL OR TRIM(e.cd_serie) = '' OR TRIM(e.cd_serie) = '-')")
            elif st in ("qualquer_incompleto", "incompleto"):
                where_clauses.append("""(
                    (e.nr_ip IS NULL OR TRIM(e.nr_ip) = '' OR TRIM(e.nr_ip) = '-')
                    OR (e.nm_hostname IS NULL OR TRIM(e.nm_hostname) = '' OR TRIM(e.nm_hostname) = '-')
                    OR (e.cd_serie IS NULL OR TRIM(e.cd_serie) = '' OR TRIM(e.cd_serie) = '-')
                )""")

        return " AND ".join(where_clauses), params

    async def criar(self, item: EquipamentoCreate, usr_id: str) -> dict:
        cd_patrimonio_clean = str(item.cd_patrimonio or "").strip().lstrip("0")
        if not cd_patrimonio_clean:
            cd_patrimonio_clean = item.cd_patrimonio.strip()

        dados_enriquecidos = await self.enriquecer_dados_equipamento(cd_patrimonio_clean)

        serie_val = item.cd_serie.strip() if item.cd_serie else dados_enriquecidos["cd_serie"]
        mac_val = item.cd_mac.strip() if item.cd_mac else dados_enriquecidos["cd_mac"]
        ip_val = item.nr_ip.strip() if item.nr_ip else dados_enriquecidos["nr_ip"]
        host_val = item.nm_hostname.strip() if item.nm_hostname else dados_enriquecidos["nm_hostname"]

        sql = """
            INSERT INTO ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS (
                cd_patrimonio,
                cd_setor,
                local_especifico,
                tipo_maquina,
                observacao,
                cd_serie,
                cd_mac,
                nr_ip,
                nm_hostname,
                registrado_por_id,
                registrado_em,
                sn_inserido_manual
            ) VALUES (
                :cd_patrimonio,
                :cd_setor,
                :local_especifico,
                :tipo_maquina,
                :observacao,
                :cd_serie,
                :cd_mac,
                :nr_ip,
                :nm_hostname,
                :registrado_por_id,
                SYSDATE,
                NVL(:sn_inserido_manual, 'N')
            )
        """
        params = {
            "cd_patrimonio": cd_patrimonio_clean,
            "cd_setor": item.cd_setor.strip(),
            "local_especifico": item.local_especifico.strip() if item.local_especifico else None,
            "tipo_maquina": item.tipo_maquina,
            "observacao": item.observacao.strip() if item.observacao else None,
            "cd_serie": serie_val,
            "cd_mac": mac_val,
            "nr_ip": ip_val,
            "nm_hostname": host_val,
            "registrado_por_id": str(usr_id),
            "sn_inserido_manual": item.sn_inserido_manual or "N",
        }
        await self.execute_dml_cimed(sql, params)
        return {"mensagem": "Equipamento cadastrado com sucesso."}

    async def criar_lote(self, payload: List[EquipamentoCreate], usr_id: str) -> dict:
        items_to_insert = []
        for item in payload:
            cd_patrimonio_clean = str(item.cd_patrimonio or "").strip().lstrip("0")
            if not cd_patrimonio_clean:
                cd_patrimonio_clean = item.cd_patrimonio.strip()

            dados_enriquecidos = await self.enriquecer_dados_equipamento(cd_patrimonio_clean)

            serie_val = item.cd_serie.strip() if item.cd_serie else dados_enriquecidos["cd_serie"]
            mac_val = item.cd_mac.strip() if item.cd_mac else dados_enriquecidos["cd_mac"]
            ip_val = item.nr_ip.strip() if item.nr_ip else dados_enriquecidos["nr_ip"]
            host_val = item.nm_hostname.strip() if item.nm_hostname else dados_enriquecidos["nm_hostname"]

            items_to_insert.append({
                "cd_patrimonio": cd_patrimonio_clean,
                "cd_setor": item.cd_setor.strip(),
                "local_especifico": item.local_especifico.strip() if item.local_especifico else None,
                "tipo_maquina": item.tipo_maquina,
                "observacao": item.observacao.strip() if item.observacao else None,
                "cd_serie": serie_val,
                "cd_mac": mac_val,
                "nr_ip": ip_val,
                "nm_hostname": host_val,
                "registrado_por_id": str(usr_id),
                "sn_inserido_manual": item.sn_inserido_manual or "N",
            })

        count = await self.execute_batch_insert_cimed(items_to_insert)
        return {"mensagem": f"{count} equipamentos cadastrados com sucesso.", "inseridos": count}

    async def consultar_patrimonio(self, cd_patrimonio: str) -> dict:
        patrimonio_limpo = str(cd_patrimonio or "").strip().lstrip("0")

        sql_setor = """
            SELECT s.cd_setor, s.nm_setor
            FROM cimed.man_detalheequipamento a
            INNER JOIN cimed.man_contratoentrega b ON a.cen_ky = b.cen_ky
            INNER JOIN cimed.man_contrato c ON b.con_ky = c.con_ky
            LEFT JOIN dbamv.setor s ON a.set_ky = s.cd_setor
            WHERE c.con_ky = 1012700
              AND a.deq_idcontrato = :cd_patrimonio
        """
        item_setor = await core.fetch_one_cimed(sql_setor, {"cd_patrimonio": patrimonio_limpo})
        dados_enriquecidos = await self.enriquecer_dados_equipamento(patrimonio_limpo)

        if not item_setor and not dados_enriquecidos["cd_serie"]:
            return None

        return {
            "cd_patrimonio": patrimonio_limpo,
            "nr_serie": dados_enriquecidos["cd_serie"],
            "cd_serie": dados_enriquecidos["cd_serie"],
            "cd_setor": item_setor.get("cd_setor") if item_setor else None,
            "nm_setor": item_setor.get("nm_setor") if item_setor else None,
            "nm_hostname": dados_enriquecidos["nm_hostname"],
            "nr_ip": dados_enriquecidos["nr_ip"],
            "mac_address": dados_enriquecidos["cd_mac"],
            "cd_mac": dados_enriquecidos["cd_mac"]
        }

    async def listar_tecnicos(self) -> List[str]:
        sql = """
            SELECT DISTINCT registrado_por_id
            FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS
            WHERE registrado_por_id IS NOT NULL
            ORDER BY registrado_por_id ASC
        """
        rows = await core.fetch_all_cimed(sql, {})
        return [r["registrado_por_id"] for r in rows if r.get("registrado_por_id")]

    async def listar_equipamentos(self, skip: int, limit: int, where_str: str, params: dict) -> dict:
        sql_count = f"""
            SELECT COUNT(*) AS total 
            FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS e 
            LEFT JOIN DBAMV.SETOR s ON s.cd_setor = e.cd_setor
            WHERE {where_str}
        """
        res_count = await core.fetch_one_cimed(sql_count, params)
        total = res_count.get("total", 0) if res_count else 0

        end_row = skip + limit
        sql_data = f"""
            SELECT cd_registro, cd_patrimonio, cd_setor, dc_setor, local_especifico,
                   tipo_maquina, observacao, cd_serie, cd_mac, nr_ip, nm_hostname, registrado_por_id, registrado_em, sn_inserido_manual
            FROM (
                SELECT cd_registro, cd_patrimonio, cd_setor, dc_setor, local_especifico,
                       tipo_maquina, observacao, cd_serie, cd_mac, nr_ip, nm_hostname, registrado_por_id,
                       TO_CHAR(registrado_em, 'YYYY-MM-DD HH24:MI:SS') AS registrado_em,
                       NVL(sn_inserido_manual, 'N') AS sn_inserido_manual,
                       ROWNUM AS rn
                FROM (
                    SELECT e.cd_registro, e.cd_patrimonio, e.cd_setor,
                           CASE WHEN s.nm_setor IS NOT NULL THEN e.cd_setor || ' - ' || s.nm_setor ELSE e.cd_setor END AS dc_setor,
                           e.local_especifico, e.tipo_maquina, e.observacao, e.cd_serie, e.cd_mac, e.nr_ip, e.nm_hostname, e.registrado_por_id, e.registrado_em, e.sn_inserido_manual
                    FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS e
                    LEFT JOIN DBAMV.SETOR s ON s.cd_setor = e.cd_setor
                    WHERE {where_str}
                    ORDER BY e.cd_registro DESC
                )
                WHERE ROWNUM <= :end_row
            )
            WHERE rn > :skip
        """
        params_data = dict(params)
        params_data["end_row"] = end_row
        params_data["skip"] = skip

        rows = await core.fetch_all_cimed(sql_data, params_data)
        return {
            "total": total,
            "skip": skip,
            "limit": limit,
            "itens": rows
        }

    async def exportar_excel(self, where_str: str, params: dict) -> io.BytesIO:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        sql = f"""
            SELECT e.cd_registro, e.cd_patrimonio, e.cd_setor,
                   CASE WHEN s.nm_setor IS NOT NULL THEN e.cd_setor || ' - ' || s.nm_setor ELSE e.cd_setor END AS dc_setor,
                   e.local_especifico, e.tipo_maquina, e.observacao, e.cd_serie, e.cd_mac, e.nr_ip, e.nm_hostname, e.registrado_por_id,
                   TO_CHAR(e.registrado_em, 'YYYY-MM-DD HH24:MI:SS') AS registrado_em,
                   NVL(e.sn_inserido_manual, 'N') AS sn_inserido_manual
            FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS e
            LEFT JOIN DBAMV.SETOR s ON s.cd_setor = e.cd_setor
            WHERE {where_str}
            ORDER BY e.cd_registro DESC
        """
        rows = await core.fetch_all_cimed(sql, params)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Equipamentos"

        headers = [
            "ID Registro", "Patrimônio", "Setor", "Local Específico", "Tipo Máquina",
            "Nº de Série", "MAC Address", "IP", "Hostname", "Registrado Por (Técnico)",
            "Data Registro", "Inserido Manual", "Observação"
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="008B95")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        data_font = Font(name="Calibri", size=10)
        data_align_left = Alignment(horizontal="left", vertical="center")
        data_align_center = Alignment(horizontal="center", vertical="center")
        zebra_fill = PatternFill(fill_type="solid", fgColor="F8FAFC")

        for row_idx, r in enumerate(rows, start=2):
            reg_em = r.get("registrado_em") or ""
            if reg_em and len(reg_em) >= 16:
                try:
                    dt_obj = datetime.strptime(reg_em[:19], "%Y-%m-%d %H:%M:%S")
                    reg_em_formatada = dt_obj.strftime("%d/%m/%Y %H:%M")
                except Exception:
                    reg_em_formatada = reg_em
            else:
                reg_em_formatada = reg_em

            row_data = [
                r.get("cd_registro"),
                r.get("cd_patrimonio") or "",
                r.get("dc_setor") or "",
                r.get("local_especifico") or "",
                r.get("tipo_maquina") or "",
                r.get("cd_serie") or "",
                r.get("cd_mac") or "",
                r.get("nr_ip") or "",
                r.get("nm_hostname") or "",
                r.get("registrado_por_id") or "",
                reg_em_formatada,
                "Sim" if r.get("sn_inserido_manual") == "S" else "Não",
                r.get("observacao") or ""
            ]
            ws.append(row_data)

            is_even = (row_idx % 2 == 0)
            for col_num in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.font = data_font
                c.border = thin_border
                if is_even:
                    c.fill = zebra_fill

                if col_num in (1, 2, 5, 6, 7, 8, 9):
                    c.alignment = data_align_center
                else:
                    c.alignment = data_align_left

            ws.row_dimensions[row_idx].height = 20

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    async def obter_devolucoes(self) -> List[dict]:
        sql = """
            SELECT 
                e.cd_registro AS CD_REGISTRO,
                e.cd_patrimonio AS CD_PATRIMONIO,
                e.cd_patrimonio AS PAT_CONTRATADA,
                e.cd_setor AS CD_SETOR,
                s.nm_setor AS NM_SETOR,
                e.local_especifico AS LOCAL_ESPECIFICO,
                e.tipo_maquina AS TIPO_MAQUINA,
                e.observacao AS OBSERVACAO,
                e.registrado_por_id AS REGISTRADO_POR_ID,
                e.registrado_em AS REGISTRADO_EM,
                NVL(e.cd_serie, e.cd_patrimonio) AS NR_SERIE,
                e.tipo_maquina AS MODELO,
                e.nr_ip AS NR_IP,
                e.cd_mac AS MAC_ADDRESS,
                'S' AS CABO_ENERGIA,
                'S' AS CABO_USB,
                'S' AS RESTAURACAO_FABRICA,
                'S' AS MOUSE,
                'S' AS TECLADO,
                'S' AS MONITOR,
                'S' AS SN_DEVOLVIDA
            FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS e
            LEFT JOIN DBAMV.SETOR s ON e.cd_setor = s.cd_setor
            ORDER BY e.cd_registro DESC
        """
        rows = await core.fetch_all_cimed(sql, {})
        if rows is None:
            rows = []

        def is_valid_sector_name(text: str) -> bool:
            if not text:
                return False
            t = text.strip().upper()
            if t.isdigit():
                return False
            if re.match(r'^[0-9A-Z]{7,15}$', t) and any(c.isdigit() for c in t) and any(c.isalpha() for c in t):
                if t != "CCIRAS":
                    return False
            return True

        sheets = {}
        sector_counters = {}

        for row in rows:
            pat = (row.get("pat_contratada") or "").strip()
            loc = (row.get("local_especifico") or "").strip()
            nm  = (row.get("nm_setor") or "").strip()

            current_sector_detected = "Complexo HCFMB (Geral)"
            if is_valid_sector_name(nm):
                current_sector_detected = nm
            elif is_valid_sector_name(loc):
                current_sector_detected = loc
            elif is_valid_sector_name(pat):
                current_sector_detected = pat

            setor_cod  = row.get("cd_setor")
            setor_completo = f"{setor_cod} - {current_sector_detected}" if setor_cod else current_sector_detected

            sec_slug = current_sector_detected.upper()
            if sec_slug not in sector_counters:
                sector_counters[sec_slug] = len(sector_counters) + 1
            
            folha_key = str(sector_counters[sec_slug])

            if folha_key not in sheets:
                data_formatada = ""
                reg_em = row.get("registrado_em")
                if reg_em:
                    if hasattr(reg_em, "strftime"):
                        data_formatada = reg_em.strftime("%d/%m/%Y")
                    else:
                        data_formatada = str(reg_em)[:10]
                if not data_formatada:
                    data_formatada = datetime.now().strftime("%d/%m/%Y")

                sheets[folha_key] = {
                    "folha": str(folha_key),
                    "setor": setor_completo,
                    "data": data_formatada,
                    "contrato": "033/2023",
                    "empresa": "COMTECH",
                    "obs": row.get("observacao") or "-",
                    "pagina": 1,
                    "equipamentos": []
                }
            
            val_sn   = "S" if (str(row.get("sn_devolvida") or "")).upper() == "S" else "N"
            val_cabo = "S" if (str(row.get("cabo_energia") or "")).upper() == "S" else "N"
            val_usb  = "S" if (str(row.get("cabo_usb") or "")).upper() == "S" else "N"
            val_rest = "S" if (str(row.get("restauracao_fabrica") or "")).upper() == "S" else "N"
            val_mou  = "S" if (str(row.get("mouse") or "")).upper() == "S" else "N"
            val_tec  = "S" if (str(row.get("teclado") or "")).upper() == "S" else "N"
            val_mon  = "S" if (str(row.get("monitor") or "")).upper() == "S" else "N"

            sheets[folha_key]["equipamentos"].append({
                "snDevolvida": val_sn,
                "nrSerie": row.get("nr_serie") or "-",
                "numEqpto": row.get("cd_patrimonio") or "-",
                "modelo": row.get("modelo") or "COMPUTADOR",
                "caboEnergia": val_cabo,
                "caboUsb": val_usb,
                "restauracao": val_rest,
                "mouse": val_mou,
                "teclado": val_tec,
                "monitor": val_mon,
                "nrIp": row.get("nr_ip") or "-",
                "macAddress": row.get("mac_address") or row.get("cd_mac") or "-",
                "localEspecifico": row.get("local_especifico") or current_sector_detected,
                "tipoMaquina": row.get("tipo_maquina") or "COMPUTADOR"
            })
            
        return list(sheets.values())

    async def obter_por_id(self, cd_registro: int) -> Optional[dict]:
        sql = """
            SELECT 
                e.cd_registro,
                e.cd_patrimonio,
                e.cd_setor,
                CASE WHEN s.nm_setor IS NOT NULL THEN e.cd_setor || ' - ' || s.nm_setor ELSE e.cd_setor END AS dc_setor,
                e.local_especifico,
                e.tipo_maquina,
                e.observacao,
                e.cd_serie,
                e.cd_mac,
                e.nr_ip,
                e.nm_hostname,
                e.registrado_por_id,
                TO_CHAR(e.registrado_em, 'YYYY-MM-DD HH24:MI:SS') AS registrado_em,
                NVL(e.sn_inserido_manual, 'N') AS sn_inserido_manual
            FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS e
            LEFT JOIN DBAMV.SETOR s ON e.cd_setor = s.cd_setor
            WHERE e.cd_registro = :cd_registro
        """
        return await core.fetch_one_cimed(sql, {"cd_registro": cd_registro})

    async def atualizar(self, cd_registro: int, payload: EquipamentoUpdate) -> int:
        fields = []
        params: Dict[str, Any] = {"cd_registro": cd_registro}

        if payload.cd_patrimonio is not None:
            fields.append("cd_patrimonio = :cd_patrimonio")
            params["cd_patrimonio"] = payload.cd_patrimonio.strip().lstrip("0") or payload.cd_patrimonio.strip()

        if payload.cd_setor is not None:
            fields.append("cd_setor = :cd_setor")
            params["cd_setor"] = payload.cd_setor.strip()

        if payload.local_especifico is not None:
            fields.append("local_especifico = :local_especifico")
            params["local_especifico"] = payload.local_especifico.strip() if payload.local_especifico else None

        if payload.tipo_maquina is not None:
            fields.append("tipo_maquina = :tipo_maquina")
            params["tipo_maquina"] = payload.tipo_maquina

        if payload.observacao is not None:
            fields.append("observacao = :observacao")
            params["observacao"] = payload.observacao.strip() if payload.observacao else None

        if payload.cd_serie is not None:
            fields.append("cd_serie = :cd_serie")
            params["cd_serie"] = payload.cd_serie.strip() if payload.cd_serie else None

        if payload.cd_mac is not None:
            fields.append("cd_mac = :cd_mac")
            params["cd_mac"] = payload.cd_mac.strip() if payload.cd_mac else None

        if payload.nr_ip is not None:
            fields.append("nr_ip = :nr_ip")
            params["nr_ip"] = payload.nr_ip.strip() if payload.nr_ip else None

        if payload.nm_hostname is not None:
            fields.append("nm_hostname = :nm_hostname")
            params["nm_hostname"] = payload.nm_hostname.strip() if payload.nm_hostname else None

        if not fields:
            return 0

        manual_val = payload.sn_inserido_manual if payload.sn_inserido_manual is not None else 'S'
        fields.append("sn_inserido_manual = :sn_inserido_manual")
        params["sn_inserido_manual"] = manual_val

        sql = f"""
            UPDATE ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS
            SET {", ".join(fields)}
            WHERE cd_registro = :cd_registro
        """
        return await self.execute_dml_cimed(sql, params)

    async def deletar(self, cd_registro: int) -> int:
        sql = "DELETE FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS WHERE cd_registro = :cd_registro"
        return await self.execute_dml_cimed(sql, {"cd_registro": cd_registro})

    async def listar_setores(self) -> List[dict]:
        sql = """
            SELECT cd_setor, nm_setor
            FROM DBAMV.SETOR
            WHERE UPPER(nm_setor) NOT LIKE '%NÃO USAR%'
              AND UPPER(nm_setor) NOT LIKE '%NAO USAR%'
            ORDER BY nm_setor ASC
        """
        rows = await core.fetch_all_cimed(sql, {})
        return [
            {
                "cd_setor": row["cd_setor"],
                "nm_setor": row["nm_setor"],
                "label": f"{row['cd_setor']} - {row['nm_setor']}"
            }
            for row in rows
        ]

    async def sincronizar_legados(self) -> dict:
        def _run():
            conn_oracle = acquire_cimed_connection()
            try:
                cur_oracle = conn_oracle.cursor()

                sql_pendentes = """
                    SELECT cd_registro, cd_patrimonio, cd_serie, nr_ip, nm_hostname, cd_mac
                    FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS
                    WHERE nr_ip IS NULL OR nm_hostname IS NULL OR cd_serie IS NULL OR cd_mac IS NULL
                """
                cur_oracle.execute(sql_pendentes)
                cols = [d[0].lower() for d in cur_oracle.description]
                pendentes = [dict(zip(cols, row)) for row in cur_oracle.fetchall()]

                total_pendentes = len(pendentes)
                if total_pendentes == 0:
                    cur_oracle.close()
                    return {
                        "total": 0,
                        "atualizados": 0,
                        "sem_dados": 0,
                        "mensagem": "Nenhum equipamento pendente de enriquecimento de dados de rede."
                    }

                sql_serie = """
                    SELECT a.deq_numserie AS cd_serie
                    FROM cimed.man_detalheequipamento a
                    INNER JOIN cimed.man_contratoentrega b ON a.cen_ky = b.cen_ky
                    INNER JOIN cimed.man_contrato c ON b.con_ky = c.con_ky
                    WHERE c.con_ky = 1012700
                      AND a.deq_idcontrato = :cd_patrimonio
                """

                sql_update = """
                    UPDATE ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS
                    SET cd_patrimonio = :cd_patrimonio_limpo,
                        cd_serie = :cd_serie,
                        cd_mac = :cd_mac,
                        nr_ip = :nr_ip,
                        nm_hostname = :nm_hostname
                    WHERE cd_registro = :cd_registro
                """

                atualizados = 0
                sem_dados = 0

                for item in pendentes:
                    cd_registro = item["cd_registro"]
                    cd_patrimonio_orig = str(item["cd_patrimonio"] or "").strip()
                    cd_patrimonio_limpo = cd_patrimonio_orig.lstrip("0")
                    if not cd_patrimonio_limpo:
                        cd_patrimonio_limpo = cd_patrimonio_orig

                    cd_serie = item.get("cd_serie")
                    if not cd_serie:
                        cur_oracle.execute(sql_serie, {"cd_patrimonio": cd_patrimonio_limpo})
                        row_serie = cur_oracle.fetchone()
                        if row_serie and row_serie[0]:
                            cd_serie = str(row_serie[0]).strip()

                    if not cd_serie:
                        sem_dados += 1
                        continue

                    nr_ip = item.get("nr_ip")
                    nm_hostname = item.get("nm_hostname")
                    cd_mac = item.get("cd_mac")

                    try:
                        conn_ocs = core.acquire_ocs_connection()
                        try:
                            with conn_ocs.cursor() as cur_ocs:
                                cur_ocs.execute("""
                                    SELECT h.NAME AS nm_hostname, h.IPADDR AS nr_ip, n.MACADDR AS cd_mac
                                    FROM hardware h
                                    INNER JOIN networks n ON h.ID = n.HARDWARE_ID
                                    INNER JOIN bios b ON h.ID = b.HARDWARE_ID
                                    WHERE b.SSN = %s
                                """, (cd_serie,))
                                row_ocs = cur_ocs.fetchone()
                                if row_ocs:
                                    if not nm_hostname:
                                        nm_hostname = row_ocs.get("nm_hostname")
                                    if not nr_ip:
                                        nr_ip = row_ocs.get("nr_ip")
                                    if not cd_mac:
                                        cd_mac = row_ocs.get("cd_mac")
                        finally:
                            conn_ocs.close()
                    except Exception as e:
                        print(f"[Equipamentos] OCS MariaDB inacessível durante sincronização para cd_serie={cd_serie}: {e}")

                    cur_oracle.execute(sql_update, {
                        "cd_patrimonio_limpo": cd_patrimonio_limpo,
                        "cd_serie": cd_serie,
                        "cd_mac": cd_mac,
                        "nr_ip": nr_ip,
                        "nm_hostname": nm_hostname,
                        "cd_registro": cd_registro
                    })
                    atualizados += 1

                conn_oracle.commit()
                cur_oracle.close()
                return {
                    "total": total_pendentes,
                    "total_processados": total_pendentes,
                    "atualizados": atualizados,
                    "total_atualizados": atualizados,
                    "sem_dados": sem_dados,
                    "mensagem": f"Sincronização concluída: {atualizados} atualizados de {total_pendentes} analisados."
                }
            except Exception as e:
                if conn_oracle:
                    conn_oracle.rollback()
                raise e
            finally:
                release_cimed_connection(conn_oracle)

        return await to_thread(_run)


class RolloutRepository:
    """Repositório responsável pelas operações de Substituição e Rollout de Equipamentos."""

    @staticmethod
    async def importar_novos(itens: List[NovoEquipamentoImport]) -> int:
        def _run():
            if not itens:
                return 0
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql_merge = """
                    MERGE INTO ESTAGIO.HCESCRITORIO_NOVO_EQUIPAMENTOS target
                    USING (
                        SELECT :cd_patrimonio AS cd_patrimonio,
                               :cd_serie AS cd_serie,
                               :cd_mac AS cd_mac,
                               :ano_rollout AS ano_rollout
                        FROM DUAL
                    ) src
                    ON (target.cd_patrimonio = src.cd_patrimonio)
                    WHEN MATCHED THEN
                        UPDATE SET target.cd_serie = NVL(src.cd_serie, target.cd_serie),
                                   target.cd_mac = NVL(src.cd_mac, target.cd_mac),
                                   target.ano_rollout = NVL(src.ano_rollout, target.ano_rollout)
                    WHEN NOT MATCHED THEN
                        INSERT (cd_patrimonio, cd_serie, cd_mac, ano_rollout, st_utilizado)
                        VALUES (src.cd_patrimonio, src.cd_serie, src.cd_mac, NVL(src.ano_rollout, 2026), 'N')
                """
                params_list = [
                    {
                        "cd_patrimonio": str(item.cd_patrimonio).strip(),
                        "cd_serie": item.cd_serie.strip() if item.cd_serie else None,
                        "cd_mac": item.cd_mac.strip() if item.cd_mac else None,
                        "ano_rollout": item.ano_rollout or 2026,
                    }
                    for item in itens
                    if item.cd_patrimonio and str(item.cd_patrimonio).strip()
                ]
                if not params_list:
                    return 0
                cur.executemany(sql_merge, params_list)
                affected = cur.rowcount
                conn.commit()
                cur.close()
                return affected
            except Exception as e:
                conn.rollback()
                raise e
            finally:
                release_cimed_connection(conn)

        return await to_thread(_run)

    @staticmethod
    async def listar_disponiveis() -> List[Dict[str, Any]]:
        def _run():
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql = """
                    SELECT cd_registro, cd_patrimonio, cd_serie, cd_mac, ano_rollout, st_utilizado, reg_key
                    FROM ESTAGIO.HCESCRITORIO_NOVO_EQUIPAMENTOS
                    WHERE st_utilizado = 'N'
                    ORDER BY cd_patrimonio ASC
                """
                cur.execute(sql)
                rows = cur.fetchall()
                cur.close()
                resultado = []
                for r in rows:
                    resultado.append({
                        "cd_registro": r[0],
                        "cd_patrimonio": r[1],
                        "cd_serie": r[2],
                        "cd_mac": r[3],
                        "ano_rollout": r[4],
                        "st_utilizado": r[5],
                        "reg_key": r[6],
                    })
                return resultado
            finally:
                release_cimed_connection(conn)

        return await to_thread(_run)

    @staticmethod
    async def executar_merge(pares: List[ParMerge]) -> int:
        def _run():
            if not pares:
                return 0
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql_update = """
                    UPDATE ESTAGIO.HCESCRITORIO_NOVO_EQUIPAMENTOS
                    SET reg_key = :reg_key_antigo,
                        st_utilizado = 'P',
                        dh_merge = SYSDATE
                    WHERE cd_registro = :cd_registro_novo
                      AND st_utilizado = 'N'
                """
                vinculados = 0
                for par in pares:
                    cur.execute(sql_update, {
                        "reg_key_antigo": par.reg_key_antigo,
                        "cd_registro_novo": par.cd_registro_novo,
                    })
                    if cur.rowcount == 0:
                        raise ValueError(f"Equipamento novo ID {par.cd_registro_novo} não está disponível para vinculação.")
                    vinculados += cur.rowcount

                conn.commit()
                cur.close()
                return vinculados
            except Exception as e:
                conn.rollback()
                raise e
            finally:
                release_cimed_connection(conn)

        return await to_thread(_run)

    @staticmethod
    async def listar_planejados(cd_setor: str) -> List[Dict[str, Any]]:
        def _run():
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql = """
                    SELECT 
                        a.cd_registro AS cd_registro_antigo,
                        a.cd_patrimonio AS cd_patrimonio_antigo,
                        a.cd_setor,
                        a.local_especifico,
                        a.tipo_maquina,
                        a.ano_rollout AS ano_rollout_antigo,
                        n.cd_registro AS cd_registro_novo,
                        n.cd_patrimonio AS cd_patrimonio_novo,
                        n.cd_serie AS cd_serie_novo,
                        n.cd_mac AS cd_mac_novo,
                        n.st_utilizado AS st_utilizado_novo
                    FROM ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS a
                    JOIN ESTAGIO.HCESCRITORIO_NOVO_EQUIPAMENTOS n ON a.cd_registro = n.reg_key
                    WHERE a.cd_setor = :cd_setor
                    ORDER BY a.local_especifico ASC, a.cd_patrimonio ASC
                """
                cur.execute(sql, {"cd_setor": str(cd_setor).strip()})
                rows = cur.fetchall()
                cur.close()
                resultado = []
                for r in rows:
                    resultado.append({
                        "cd_registro_antigo": r[0],
                        "cd_patrimonio_antigo": r[1],
                        "cd_setor": r[2],
                        "local_especifico": r[3],
                        "tipo_maquina": r[4],
                        "ano_rollout_antigo": r[5],
                        "cd_registro_novo": r[6],
                        "cd_patrimonio_novo": r[7],
                        "cd_serie_novo": r[8],
                        "cd_mac_novo": r[9],
                        "st_utilizado_novo": r[10],
                    })
                return resultado
            finally:
                release_cimed_connection(conn)

        return await to_thread(_run)

    @staticmethod
    async def confirmar_devolucao(cd_registro_antigo: int) -> bool:
        def _run():
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql = """
                    UPDATE ESTAGIO.HCESCRITORIO_REG_EQUIPAMENTOS
                    SET ano_rollout = TO_NUMBER(TO_CHAR(SYSDATE, 'YYYY'))
                    WHERE cd_registro = :cd_registro_antigo
                """
                cur.execute(sql, {"cd_registro_antigo": cd_registro_antigo})
                affected = cur.rowcount
                conn.commit()
                cur.close()
                return affected > 0
            except Exception as e:
                conn.rollback()
                raise e
            finally:
                release_cimed_connection(conn)

        return await to_thread(_run)

    @staticmethod
    async def confirmar_entrega(cd_registro_novo: int) -> bool:
        def _run():
            conn = acquire_cimed_connection()
            try:
                cur = conn.cursor()
                sql = """
                    UPDATE ESTAGIO.HCESCRITORIO_NOVO_EQUIPAMENTOS
                    SET st_utilizado = 'S',
                        dt_substituicao = SYSDATE
                    WHERE cd_registro = :cd_registro_novo
                """
                cur.execute(sql, {"cd_registro_novo": cd_registro_novo})
                affected = cur.rowcount
                conn.commit()
                cur.close()
                return affected > 0
            except Exception as e:
                conn.rollback()
                raise e
            finally:
                release_cimed_connection(conn)

        return await to_thread(_run)

